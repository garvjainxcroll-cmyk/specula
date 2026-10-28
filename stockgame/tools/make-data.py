"""Make stockgame/secret/game.json from the organiser's two sources.

  python3 -I stockgame/tools/make-data.py <price-impact-sheet.xlsx> <updates.txt>

- The price-impact sheet (Excel, "Impact Log" and "Stocks" tabs) gives the ten
  stocks and the 180 price moves (45 updates x 4 stocks), each with its reason.
- updates.txt is the participant text of the 45 updates ("Update N · Category",
  the headline on the next line, then one paragraph per line).

Release codes are made once and kept: running this again keeps the codes and
salt already in game.json, so a rebuilt website still accepts the same codes.
game.json is secret (it holds every future headline, move and code) and is
never committed: the repository is public.
"""
import json
import os
import re
import secrets
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, '..', 'secret', 'game.json')
ALPHABET = 'ABCDEFGHJKMNPQRSTUVWXYZ23456789'  # no 0/O, 1/I/L


def read_sheet(path):
    wb = openpyxl.load_workbook(path)
    stocks = []
    for r in wb['Stocks'].iter_rows(min_row=5, values_only=True):
        if not isinstance(r[0], int):
            continue
        stocks.append({'t': r[2], 'name': r[1], 'sector': r[3], 'start': r[4]})
    moves = []
    for r in wb['Impact Log'].iter_rows(min_row=5, values_only=True):
        if r[0] is None:
            continue
        moves.append({'u': int(r[0]), 'cat': r[2], 'head': r[3], 't': r[4], 'pct': float(r[7]), 'why': r[10]})
    return stocks, moves


def read_updates(path):
    updates, cur = [], None
    for line in open(path, encoding='utf-8'):
        line = line.strip()
        if not line or line.startswith('Round '):
            continue
        m = re.match(r'^Update (\d+) · (.+)$', line)
        if m:
            cur = {'i': int(m.group(1)), 'cat': m.group(2), 'head': None, 'body': []}
            updates.append(cur)
        elif cur['head'] is None:
            cur['head'] = line
        else:
            cur['body'].append(line)
    return updates


def new_code(used):
    while True:
        c = ''.join(secrets.choice(ALPHABET) for _ in range(8))
        if c not in used:
            used.add(c)
            return c


def main():
    sheet, text = sys.argv[1], sys.argv[2]
    stocks, moves = read_sheet(sheet)
    updates = read_updates(text)
    assert len(stocks) == 10, len(stocks)
    assert [u['i'] for u in updates] == list(range(1, 46)), 'updates must be 1..45 in order'
    tickers = {s['t'] for s in stocks}
    problems = []
    for u in updates:
        mine = [m for m in moves if m['u'] == u['i']]
        if len(mine) != 4:
            problems.append(f"update {u['i']}: {len(mine)} moves, expected 4")
        for m in mine:
            if m['t'] not in tickers:
                problems.append(f"update {u['i']}: unknown ticker {m['t']}")
            if m['cat'] != u['cat']:
                problems.append(f"update {u['i']}: category '{u['cat']}' in text, '{m['cat']}' in sheet")
            if m['head'] != u['head']:
                problems.append(f"update {u['i']}: headline differs\n  text:  {u['head']}\n  sheet: {m['head']}")
        u['round'] = (u['i'] + 4) // 5
        u['moves'] = [{'t': m['t'], 'pct': m['pct'], 'why': m['why']} for m in mine]
    if problems:
        print('\n'.join(problems))
        sys.exit(1)

    old = json.load(open(OUT, encoding='utf-8')) if os.path.exists(OUT) else None
    if old and 'codes' in old:
        codes, salt = old['codes'], old['salt']
        print('Kept the existing release codes and salt.')
    else:
        used = set()
        codes = {'N': [new_code(used) for _ in range(45)], 'P': [new_code(used) for _ in range(45)]}
        salt = secrets.token_hex(16)
        print('Made new release codes.')

    game = {
        'v': 1,
        'note': 'SECRET. Every future headline, price move and release code. Never publish or commit.',
        'capital': 1000000,
        'stocks': stocks,
        'updates': updates,
        'codes': codes,
        'salt': salt,
        'iterations': 100000,
    }
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, 'w', encoding='utf-8') as f:
        json.dump(game, f, ensure_ascii=False, indent=1)
    print(f'Wrote {os.path.normpath(OUT)}: {len(stocks)} stocks, {len(updates)} updates, {sum(len(u["moves"]) for u in updates)} moves.')


if __name__ == '__main__':
    main()
